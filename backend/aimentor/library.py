import hashlib

import csv

import io

import json

import sqlite3

import time

import uuid

import zipfile

import openpyxl

from fastapi import HTTPException

from .database import dump

from .common import error, iso, ALL_SUBJECTS, ACTIVE_SUBJECTS



SUBJECT_ALIASES = {

    "math": "math", "toan": "math", "toán": "math",

    "literature": "literature", "van": "literature", "văn": "literature", "ngu van": "literature", "ngữ văn": "literature",

    "english": "english", "tieng anh": "english", "tiếng anh": "english", "anh": "english",

    "history": "history", "su": "history", "sử": "history", "lich su": "history", "lịch sử": "history",

    "geography": "geography", "dia": "geography", "địa": "geography", "dia ly": "geography", "địa lý": "geography",

    "chemistry": "chemistry", "hoa": "chemistry", "hóa": "chemistry", "hoa hoc": "chemistry", "hóa học": "chemistry",

    "physics": "physics", "ly": "physics", "lý": "physics", "vat ly": "physics", "vật lý": "physics",

    "biology": "biology", "sinh": "biology", "sinh hoc": "biology", "sinh học": "biology",

    "economics_law": "economics_law", "economics-law": "economics_law", "gdkt": "economics_law", "gdkt & pl": "economics_law", "gdkt&pl": "economics_law"

}



COLUMN_MAP = {

    "mã câu": "code", "mã": "code", "code": "code", "id": "code",

    "môn (wireid)": "subjectId", "môn": "subjectId", "subject": "subjectId", "subjectid": "subjectId", "subject_id": "subjectId",

    "lớp (10-12)": "grade", "lớp": "grade", "grade": "grade",

    "chương": "chapter", "chapter": "chapter",

    "chủ đề": "topic", "topic": "topic",

    "bộ sách": "textbook", "textbook": "textbook",

    "câu hỏi": "prompt", "nội dung": "prompt", "prompt": "prompt", "question": "prompt",

    "đáp án a": "optionA", "a": "optionA", "option_a": "optionA", "optiona": "optionA",

    "đáp án b": "optionB", "b": "optionB", "option_b": "optionB", "optionb": "optionB",

    "đáp án c": "optionC", "c": "optionC", "option_c": "optionC", "optionc": "optionC",

    "đáp án d": "optionD", "d": "optionD", "option_d": "optionD", "optiond": "optionD",

    "đáp án đúng (a/b/c/d)": "correctOption", "đáp án đúng": "correctOption", "correct_option": "correctOption", "correctoption": "correctOption", "answer": "correctOption",

    "lời giải": "explanation", "giải thích": "explanation", "explanation": "explanation",

    "độ khó (easy/medium/hard)": "difficulty", "độ khó": "difficulty", "difficulty": "difficulty",

    "điểm": "points", "points": "points",

    "nguồn": "source", "source": "source",

    "tác giả": "author", "author": "author",

    "quyền": "license", "bản quyền": "license", "license": "license"

}



REQUIRED_FIELD_LABELS = {

    "prompt": "Câu hỏi",

    "optionA": "Đáp án A",

    "optionB": "Đáp án B",

    "optionC": "Đáp án C",

    "optionD": "Đáp án D",

    "correctOption": "Đáp án đúng"

}



def validate_single_question_dict(row: dict, row_idx: int = 1, seen_codes: set = None, existing_db_codes: set = None):

    if not isinstance(row, dict):

        return None, f"Dòng {row_idx}: Dữ liệu không đúng định dạng bản ghi (phải là dictionary)"



    prompt = str(row.get("prompt") or "").strip()

    if not prompt:

        return None, f"Dòng {row_idx}: Nội dung câu hỏi không được để trống"

    if len(prompt) > 5000:

        return None, f"Dòng {row_idx}: Nội dung câu hỏi (prompt) vượt quá độ dài tối đa (5000 ký tự)"



    option_a = str(row.get("optionA") or "").strip()

    option_b = str(row.get("optionB") or "").strip()

    option_c = str(row.get("optionC") or "").strip()

    option_d = str(row.get("optionD") or "").strip()

    if not option_a or not option_b or not option_c or not option_d:

        return None, f"Dòng {row_idx}: Đủ 4 phương án A, B, C, D không được để trống"

    if len(option_a) > 2000:

        return None, f"Dòng {row_idx}: Phương án A vượt quá độ dài tối đa (2000 ký tự)"

    if len(option_b) > 2000:

        return None, f"Dòng {row_idx}: Phương án B vượt quá độ dài tối đa (2000 ký tự)"

    if len(option_c) > 2000:

        return None, f"Dòng {row_idx}: Phương án C vượt quá độ dài tối đa (2000 ký tự)"

    if len(option_d) > 2000:

        return None, f"Dòng {row_idx}: Phương án D vượt quá độ dài tối đa (2000 ký tự)"



    raw_correct = row.get("correctOption")

    if raw_correct is None or not str(raw_correct).strip():

        return None, f"Dòng {row_idx}: Đáp án đúng không được để trống (phải là A, B, C hoặc D)"

    correct_option = str(raw_correct).strip().upper()

    if correct_option not in ("A", "B", "C", "D"):

        return None, f"Dòng {row_idx}: Đáp án đúng phải là A, B, C hoặc D (nhận được: '{raw_correct}')"



    raw_sub = str(row.get("subjectId") or "physics").strip().lower()

    subject_id = SUBJECT_ALIASES.get(raw_sub, raw_sub)

    if subject_id not in ALL_SUBJECTS:

        return None, f"Dòng {row_idx}: Môn học '{raw_sub}' không hợp lệ. Môn hợp lệ: {', '.join(sorted(ALL_SUBJECTS))}"



    raw_grade = row.get("grade")

    if raw_grade is None or str(raw_grade).strip() == "":

        grade = 10

    else:

        try:

            grade = int(str(raw_grade).strip())

        except ValueError:

            return None, f"Dòng {row_idx}: Khối lớp phải là 10, 11 hoặc 12 (nhận được: '{raw_grade}')"

    if grade not in (10, 11, 12):

        return None, f"Dòng {row_idx}: Khối lớp phải là 10, 11 hoặc 12 (nhận được: '{grade}')"



    raw_diff = str(row.get("difficulty") or "MEDIUM").strip().upper()

    if raw_diff not in ("EASY", "MEDIUM", "HARD"):

        return None, f"Dòng {row_idx}: Độ khó phải là EASY, MEDIUM hoặc HARD (nhận được: '{row.get('difficulty')}')"

    difficulty = raw_diff



    raw_points = row.get("points")

    if raw_points is None or str(raw_points).strip() == "":

        points = 10

    else:

        try:

            points = int(str(raw_points).strip())

        except ValueError:

            return None, f"Dòng {row_idx}: Điểm câu hỏi phải là số nguyên (nhận được: '{raw_points}')"

    if points < 1 or points > 100:

        return None, f"Dòng {row_idx}: Điểm câu hỏi phải nằm trong khoảng từ 1 đến 100 (nhận được: {points})"



    raw_code = str(row.get("code") or "").strip()

    code = raw_code or f"Q_IMP_{uuid.uuid4().hex[:8].upper()}"

    if len(code) > 64:

        return None, f"Dòng {row_idx}: Mã câu hỏi '{code}' vượt quá 64 ký tự"



    if seen_codes is not None:

        if code in seen_codes:

            return None, f"Dòng {row_idx}: Mã câu hỏi '{code}' bị trùng lặp trong file import"

        seen_codes.add(code)



    if existing_db_codes is not None and code in existing_db_codes:

        return None, f"Dòng {row_idx}: Mã câu hỏi '{code}' đã tồn tại trong hệ thống. Không thể import đè."



    raw_explanation = str(row.get("explanation") or "").strip()

    if len(raw_explanation) > 5000:

        return None, f"Dòng {row_idx}: Lời giải (explanation) vượt quá độ dài tối đa (5000 ký tự)"

    explanation = raw_explanation



    raw_chapter = str(row.get("chapter") or "").strip()

    if len(raw_chapter) > 200:

        return None, f"Dòng {row_idx}: Tên chương (chapter) vượt quá độ dài tối đa (200 ký tự)"

    chapter = raw_chapter if raw_chapter else "Tổng hợp"



    raw_topic = str(row.get("topic") or "").strip()

    if len(raw_topic) > 200:

        return None, f"Dòng {row_idx}: Chủ đề (topic) vượt quá độ dài tối đa (200 ký tự)"

    topic = raw_topic if raw_topic else "Chủ đề chung"



    raw_textbook = str(row.get("textbook") or "").strip()

    if len(raw_textbook) > 100:

        return None, f"Dòng {row_idx}: Bộ sách (textbook) vượt quá độ dài tối đa (100 ký tự)"

    textbook = raw_textbook if raw_textbook else "Chưa xác định"



    raw_source = str(row.get("source") or "").strip()

    if len(raw_source) > 200:

        return None, f"Dòng {row_idx}: Nguồn trích dẫn (source) vượt quá độ dài tối đa (200 ký tự)"

    source = raw_source if raw_source else "Giáo viên upload"



    raw_author = str(row.get("author") or "").strip()

    if len(raw_author) > 100:

        return None, f"Dòng {row_idx}: Tác giả (author) vượt quá độ dài tối đa (100 ký tự)"

    author = raw_author if raw_author else "Người biên soạn"



    raw_license = str(row.get("license") or "").strip()

    if len(raw_license) > 100:

        return None, f"Dòng {row_idx}: Bản quyền (license) vượt quá độ dài tối đa (100 ký tự)"

    license_str = raw_license if raw_license else "Nội bộ"



    cleaned = {

        "code": code,

        "subjectId": subject_id,

        "grade": grade,

        "chapter": chapter,

        "topic": topic,

        "textbook": textbook,

        "prompt": prompt,

        "optionA": option_a,

        "optionB": option_b,

        "optionC": option_c,

        "optionD": option_d,

        "correctOption": correct_option,

        "explanation": explanation,

        "difficulty": difficulty,

        "points": points,

        "source": source,

        "author": author,

        "license": license_str,

        "status": "PUBLISHED" if str(row.get("status") or "").upper() == "PUBLISHED" else "DRAFT"

    }

    return cleaned, None





class LibraryService:

    def __init__(self, database):

        self.db = database



    def seed_sample_physics(self):

        with self.db.transaction() as db:

            existing = db.execute("SELECT COUNT(*) FROM questions WHERE subject_id='physics' AND grade=10").fetchone()[0]

            if existing > 0:

                return

            now = time.time()

            sample_questions = [

                {

                    "id": "q_phy10_01",

                    "code": "PHY10_CH1_Q01",

                    "subject_id": "physics",

                    "grade": 10,

                    "chapter": "Chương 1: Mở đầu & Động học",

                    "topic": "Chuyển động thẳng đều",

                    "textbook": "Kết nối tri thức",

                    "prompt": "Phương trình chuyển động của một chất điểm dọc theo trục Ox có dạng $x = 5 + 6t$ (m; s). Vận tốc của chất điểm là:",

                    "option_a": "5 m/s",

                    "option_b": "6 m/s",

                    "option_c": "-5 m/s",

                    "option_d": "11 m/s",

                    "correct_option": "B",

                    "explanation": "So sánh với phương trình tổng quát $x = x_0 + vt$, ta thấy vận tốc $v = 6\\text{ m/s}$.",

                    "difficulty": "EASY",

                    "points": 10,

                    "source": "Bộ đề ôn tập Vật lý 10",

                    "author": "Ban biên soạn AI Mentor",

                    "license": "Nội bộ",

                    "status": "PUBLISHED"

                },

                {

                    "id": "q_phy10_02",

                    "code": "PHY10_CH1_Q02",

                    "subject_id": "physics",

                    "grade": 10,

                    "chapter": "Chương 1: Mở đầu & Động học",

                    "topic": "Chuyển động thẳng biến đổi đều",

                    "textbook": "Kết nối tri thức",

                    "prompt": "Công thức liên hệ giữa vận tốc $v$, gia tốc $a$ và độ dịch chuyển $s$ trong chuyển động thẳng biến đổi đều là:",

                    "option_a": "$v^2 - v_0^2 = 2as$",

                    "option_b": "$v - v_0 = 2as$",

                    "option_c": "$v^2 + v_0^2 = 2as$",

                    "option_d": "$v^2 - v_0^2 = as$",

                    "correct_option": "A",

                    "explanation": "Công thức độc lập thời gian là $v^2 - v_0^2 = 2as$.",

                    "difficulty": "EASY",

                    "points": 10,

                    "source": "Bộ đề ôn tập Vật lý 10",

                    "author": "Ban biên soạn AI Mentor",

                    "license": "Nội bộ",

                    "status": "PUBLISHED"

                },

                {

                    "id": "q_phy10_03",

                    "code": "PHY10_CH2_Q03",

                    "subject_id": "physics",

                    "grade": 10,

                    "chapter": "Chương 2: Động lực học",

                    "topic": "Định luật II Newton",

                    "textbook": "Cánh Diều",

                    "prompt": "Một lực $F = 10\\text{ N}$ tác dụng vào vật khối lượng $m = 2\\text{ kg}$. Gia tốc của vật là:",

                    "option_a": "$5\\text{ m/s}^2$",

                    "option_b": "$20\\text{ m/s}^2$",

                    "option_c": "$0.2\\text{ m/s}^2$",

                    "option_d": "$8\\text{ m/s}^2$",

                    "correct_option": "A",

                    "explanation": "Theo Định luật II Newton: $a = \\frac{F}{m} = \\frac{10}{2} = 5\\text{ m/s}^2$.",

                    "difficulty": "MEDIUM",

                    "points": 10,

                    "source": "Bộ đề ôn tập Vật lý 10",

                    "author": "Ban biên soạn AI Mentor",

                    "license": "Nội bộ",

                    "status": "PUBLISHED"

                },

                {

                    "id": "q_phy10_04",

                    "code": "PHY10_CH1_Q04",

                    "subject_id": "physics",

                    "grade": 10,

                    "chapter": "Chương 1: Mở đầu & Động học",

                    "topic": "Sự rơi tự do",

                    "textbook": "Chân trời sáng tạo",

                    "prompt": "Chuyển động rơi tự do của một vật là chuyển động:",

                    "option_a": "Thẳng đều",

                    "option_b": "Thẳng nhanh dần đều với gia tốc g",

                    "option_c": "Thẳng chậm dần đều",

                    "option_d": "Cong đều",

                    "correct_option": "B",

                    "explanation": "Rơi tự do là chuyển động thẳng nhanh dần đều dưới tác dụng duy nhất của trọng lực với gia tốc rơi tự do g.",

                    "difficulty": "EASY",

                    "points": 10,

                    "source": "Bộ đề ôn tập Vật lý 10",

                    "author": "Ban biên soạn AI Mentor",

                    "license": "Nội bộ",

                    "status": "PUBLISHED"

                },

                {

                    "id": "q_phy10_05",

                    "code": "PHY10_CH2_Q05",

                    "subject_id": "physics",

                    "grade": 10,

                    "chapter": "Chương 2: Động lực học",

                    "topic": "Tổng hợp và phân tích lực",

                    "textbook": "Kết nối tri thức",

                    "prompt": "Hai lực đồng quy $F_1 = 3\\text{ N}$ và $F_2 = 4\\text{ N}$ vuông góc với nhau. Độ lớn của hợp lực $F$ là:",

                    "option_a": "1 N",

                    "option_b": "7 N",

                    "option_c": "5 N",

                    "option_d": "12 N",

                    "correct_option": "C",

                    "explanation": "Vì hai lực vuông góc nên $F = \\sqrt{F_1^2 + F_2^2} = \\sqrt{3^2 + 4^2} = 5\\text{ N}$.",

                    "difficulty": "MEDIUM",

                    "points": 10,

                    "source": "Bộ đề ôn tập Vật lý 10",

                    "author": "Ban biên soạn AI Mentor",

                    "license": "Nội bộ",

                    "status": "PUBLISHED"

                }

            ]



            for q in sample_questions:

                db.execute("""

                INSERT OR REPLACE INTO questions

                (id, code, version, subject_id, grade, chapter, topic, textbook, prompt, option_a, option_b, option_c, option_d, correct_option, explanation, difficulty, points, source, author, license, status, created_at, updated_at)

                VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)

                """, (q["id"], q["code"], q["subject_id"], q["grade"], q["chapter"], q["topic"], q["textbook"],

                      q["prompt"], q["option_a"], q["option_b"], q["option_c"], q["option_d"], q["correct_option"],

                      q["explanation"], q["difficulty"], q["points"], q["source"], q["author"], q["license"],

                      q["status"], now, now))



            quiz_id = "quiz_phy10_basic"

            q_ids = json.dumps([q["id"] for q in sample_questions])

            db.execute("""

            INSERT OR REPLACE INTO library_quizzes (id, title, subject_id, grade, chapter, topic, difficulty, question_ids, status, created_at)

            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)

            """, (quiz_id, "Bộ đề ôn tập Vật lý 10 - Động học & Động lực học", "physics", 10, "Tổng hợp", "Tổng hợp Lớp 10", "MEDIUM", q_ids, "PUBLISHED", now))



    def list_quizzes(self, subject_id, grade, difficulty=None, chapter=None, topic=None):

        norm_sub = SUBJECT_ALIASES.get(subject_id.strip().lower(), subject_id.strip().lower())

        query = "SELECT * FROM library_quizzes WHERE subject_id=? AND grade=? AND status='PUBLISHED'"

        params = [norm_sub, grade]

        if difficulty and difficulty.upper() != "ALL":

            query += " AND difficulty=?"

            params.append(difficulty.upper())

        if chapter:

            query += " AND chapter LIKE ?"

            params.append(f"%{chapter}%")

        if topic:

            query += " AND topic LIKE ?"

            params.append(f"%{topic}%")

        query += " ORDER BY created_at DESC"



        with self.db.read() as db:

            rows = db.execute(query, params).fetchall()

            quizzes = []

            for r in rows:

                try:

                    q_ids = json.loads(r["question_ids"])

                except Exception:

                    q_ids = []

                if not q_ids:

                    continue



                # Consistent policy: Quiz must have ALL questions PUBLISHED and matching subject/grade

                placeholders = ",".join("?" * len(q_ids))

                q_rows = db.execute(f"SELECT id, status, subject_id, grade FROM questions WHERE id IN ({placeholders})", q_ids).fetchall()

                q_map = {row["id"]: row for row in q_rows}

                if len(q_map) != len(q_ids):

                    continue

                if any(row["status"] != "PUBLISHED" for row in q_map.values()):

                    continue

                if any(SUBJECT_ALIASES.get(row["subject_id"].lower(), row["subject_id"].lower()) != norm_sub or row["grade"] != grade for row in q_map.values()):

                    continue



                quizzes.append({

                    "id": r["id"],

                    "title": r["title"],

                    "subjectId": r["subject_id"],

                    "grade": r["grade"],

                    "chapter": r["chapter"],

                    "topic": r["topic"],

                    "difficulty": r["difficulty"],

                    "questionIds": q_ids,

                    "questionCount": len(q_ids),

                    "createdAt": iso(r["created_at"])

                })

            return quizzes



    def get_quiz_for_student(self, quiz_id):

        with self.db.read() as db:

            quiz = db.execute("SELECT * FROM library_quizzes WHERE id=? AND status='PUBLISHED'", (quiz_id,)).fetchone()

            if not quiz:

                error("QUIZ_NOT_FOUND", 404)

            try:

                q_ids = json.loads(quiz["question_ids"])

            except Exception:

                q_ids = []

            if not q_ids:

                error("QUIZ_NOT_FOUND", 404)



            placeholders = ",".join("?" * len(q_ids))

            q_rows = db.execute(f"SELECT * FROM questions WHERE id IN ({placeholders}) AND status='PUBLISHED'", q_ids).fetchall()

            q_map = {r["id"]: r for r in q_rows}



            # Check that ALL question IDs exist, are PUBLISHED, and match subject & grade

            quiz_sub = SUBJECT_ALIASES.get(quiz["subject_id"].lower(), quiz["subject_id"].lower())

            for qid in q_ids:

                if qid not in q_map:

                    error("QUIZ_NOT_FOUND", 404)

                r = q_map[qid]

                r_sub = SUBJECT_ALIASES.get(r["subject_id"].lower(), r["subject_id"].lower())

                if r_sub != quiz_sub or r["grade"] != quiz["grade"]:

                    error("QUIZ_NOT_FOUND", 404)



            questions = []

            for qid in q_ids:

                r = q_map[qid]

                questions.append({

                    "id": r["id"],

                    "code": r["code"],

                    "chapter": r["chapter"],

                    "topic": r["topic"],

                    "prompt": r["prompt"],

                    "options": [

                        {"id": "A", "text": r["option_a"]},

                        {"id": "B", "text": r["option_b"]},

                        {"id": "C", "text": r["option_c"]},

                        {"id": "D", "text": r["option_d"]}

                    ],

                    "difficulty": r["difficulty"],

                    "points": r["points"],

                    "imageUrl": r["image_url"]

                })

            return {

                "id": quiz["id"],

                "title": quiz["title"],

                "subjectId": quiz["subject_id"],

                "grade": quiz["grade"],

                "chapter": quiz["chapter"],

                "topic": quiz["topic"],

                "difficulty": quiz["difficulty"],

                "questions": questions

            }



    def start_session(self, user_id, quiz_id):

        with self.db.read() as db:

            quiz = db.execute("SELECT * FROM library_quizzes WHERE id=? AND status='PUBLISHED'", (quiz_id,)).fetchone()

            if not quiz:

                error("QUIZ_NOT_FOUND", 404)

            try:

                q_ids = json.loads(quiz["question_ids"])

            except Exception:

                q_ids = []

            if not q_ids:

                error("QUIZ_NOT_AVAILABLE", 400)



            placeholders = ",".join("?" * len(q_ids))

            q_rows = db.execute(f"SELECT * FROM questions WHERE id IN ({placeholders}) AND status='PUBLISHED'", q_ids).fetchall()

            q_map = {r["id"]: r for r in q_rows}



            # Consistent check: Cannot start session if any question is missing, draft, or mismatched

            quiz_sub = SUBJECT_ALIASES.get(quiz["subject_id"].lower(), quiz["subject_id"].lower())

            for qid in q_ids:

                if qid not in q_map:

                    raise HTTPException(400, detail={"code": "QUIZ_NOT_AVAILABLE", "message": f"Bộ đề không thể sử dụng vì câu hỏi '{qid}' chưa xuất bản hoặc đã bị xóa"})

                r = q_map[qid]

                r_sub = SUBJECT_ALIASES.get(r["subject_id"].lower(), r["subject_id"].lower())

                if r_sub != quiz_sub or r["grade"] != quiz["grade"]:

                    raise HTTPException(400, detail={"code": "QUIZ_NOT_AVAILABLE", "message": f"Bộ đề không thể sử dụng vì câu hỏi '{r['code'] or qid}' không cùng môn/lớp"})



            snapshots = []

            for qid in q_ids:

                r = q_map[qid]

                snapshots.append({

                    "id": r["id"],

                    "code": r["code"],

                    "version": r["version"],

                    "chapter": r["chapter"],

                    "topic": r["topic"],

                    "prompt": r["prompt"],

                    "options": [

                        {"id": "A", "text": r["option_a"]},

                        {"id": "B", "text": r["option_b"]},

                        {"id": "C", "text": r["option_c"]},

                        {"id": "D", "text": r["option_d"]}

                    ],

                    "difficulty": r["difficulty"],

                    "points": r["points"],

                    "imageUrl": r["image_url"],

                    "correctOption": r["correct_option"],

                    "explanation": r["explanation"]

                })



        session_id = f"psess_{uuid.uuid4().hex}"

        now = time.time()

        snapshots_json = json.dumps(snapshots, ensure_ascii=False)



        with self.db.transaction() as db:

            db.execute("""

            INSERT INTO practice_sessions (id, user_id, quiz_id, subject_id, grade, question_snapshots, step, answers, state, score, correct_count, total_count, created_at, updated_at)

            VALUES (?, ?, ?, ?, ?, ?, 0, '{}', 'IN_PROGRESS', 0, 0, ?, ?, ?)

            """, (session_id, user_id, quiz_id, quiz["subject_id"], quiz["grade"], snapshots_json, len(snapshots), now, now))



        return self.get_session(user_id, session_id)



    def get_session(self, user_id, session_id):

        with self.db.read() as db:

            row = db.execute("SELECT * FROM practice_sessions WHERE id=? AND user_id=?", (session_id, user_id)).fetchone()

            if not row:

                error("SESSION_NOT_FOUND", 404)



            questions = json.loads(row["question_snapshots"])

            answers = json.loads(row["answers"])

            max_score = sum(q.get("points", 10) for q in questions)



            enriched_questions = []

            for q in questions:

                item = {

                    "id": q["id"],

                    "code": q.get("code"),

                    "chapter": q.get("chapter"),

                    "topic": q.get("topic"),

                    "prompt": q["prompt"],

                    "options": q["options"],

                    "difficulty": q.get("difficulty"),

                    "points": q.get("points"),

                    "imageUrl": q.get("imageUrl")

                }

                ans = answers.get(q["id"])

                show_solution = (ans is not None) or (row["state"] == "COMPLETED")

                if ans:

                    item["userAnswer"] = ans.get("selectedOption")

                    item["isCorrect"] = ans.get("isCorrect")

                if show_solution:

                    if "correctOption" in q:

                        item["correctOption"] = q["correctOption"]

                        item["explanation"] = q.get("explanation", "")

                    else:

                        # Legacy session lacking snapshot data:

                        # NEVER query current questions table for mutated answers!

                        item["correctOption"] = None

                        item["explanation"] = "Dữ liệu lời giải lịch sử không khả dụng cho phiên làm bài cũ."

                enriched_questions.append(item)



            first_unanswered_idx = 0

            for idx, q in enumerate(questions):

                if q["id"] not in answers:

                    first_unanswered_idx = idx

                    break

            else:

                first_unanswered_idx = max(0, len(questions) - 1)



            return {

                "id": row["id"],

                "quizId": row["quiz_id"],

                "subjectId": row["subject_id"],

                "grade": row["grade"],

                "step": row["step"],

                "state": row["state"],

                "score": row["score"],

                "maxScore": max_score,

                "correctCount": row["correct_count"],

                "totalCount": row["total_count"],

                "answeredCount": len(answers),

                "firstUnansweredIndex": first_unanswered_idx,

                "questions": enriched_questions,

                "answers": answers,

                "createdAt": iso(row["created_at"]),

                "updatedAt": iso(row["updated_at"])

            }



    def submit_answer(self, user_id, session_id, question_id, selected_option):

        if not isinstance(selected_option, str):

            raise HTTPException(400, detail={"code": "INVALID_OPTION", "message": "Phương án chọn phải là chuỗi A, B, C hoặc D."})

        selected_clean = selected_option.strip().upper()

        if selected_clean not in ("A", "B", "C", "D"):

            raise HTTPException(400, detail={"code": "INVALID_OPTION", "message": "Phương án chọn phải là A, B, C hoặc D."})



        with self.db.transaction() as db:

            session = db.execute("SELECT * FROM practice_sessions WHERE id=? AND user_id=?", (session_id, user_id)).fetchone()

            if not session:

                error("SESSION_NOT_FOUND", 404)

            if session["state"] == "COMPLETED":

                error("SESSION_ALREADY_COMPLETED", 409)



            answers = json.loads(session["answers"])

            if question_id in answers:

                # Idempotency guard: if already answered, retain first answer and do not double count

                return self.get_session(user_id, session_id)



            try:

                snapshots = json.loads(session["question_snapshots"])

            except Exception:

                snapshots = []

            target_q = next((q for q in snapshots if q["id"] == question_id), None)

            if not target_q:

                raise HTTPException(404, detail={"code": "QUESTION_NOT_FOUND", "message": f"Câu hỏi '{question_id}' không thuộc bài luyện này."})



            if "correctOption" in target_q:

                correct_option = target_q["correctOption"]

                points = target_q.get("points", 10)

                explanation = target_q.get("explanation", "")

            else:

                # In-progress legacy session lacking snapshot data

                # Cannot grade against mutated current questions table!

                raise HTTPException(400, detail={

                    "code": "LEGACY_SESSION_REQUIRES_RESTART",

                    "message": "Phiên làm bài cũ thiếu dữ liệu snapshot để chấm điểm chính xác. Vui lòng bắt đầu phiên luyện tập mới."

                })



            correct = (selected_clean == correct_option.upper().strip())

            now = time.time()

            answers[question_id] = {

                "selectedOption": selected_clean,

                "isCorrect": correct,

                "answeredAt": now

            }



            # Recalculate score and correct_count deterministically to prevent drift

            correct_count = sum(1 for qid, a in answers.items() if a.get("isCorrect"))

            new_score = sum(

                next((q.get("points", 10) for q in snapshots if q["id"] == qid), 10)

                for qid, a in answers.items() if a.get("isCorrect")

            )



            db.execute("""

            UPDATE practice_sessions SET answers=?, score=?, correct_count=?, updated_at=? WHERE id=?

            """, (json.dumps(answers, ensure_ascii=False), new_score, correct_count, now, session_id))



            if not correct:

                snapshot_json = json.dumps(target_q, ensure_ascii=False)

                new_id = f"wq_{uuid.uuid4().hex[:8]}"

                db.execute("""

                INSERT INTO wrong_questions_notebook (id, user_id, question_id, subject_id, grade, selected_option, correct_option, explanation, question_snapshot, state, created_at, updated_at)

                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'NEEDS_REVIEW', ?, ?)

                ON CONFLICT(user_id, question_id) DO UPDATE SET

                  selected_option=excluded.selected_option,

                  correct_option=excluded.correct_option,

                  explanation=excluded.explanation,

                  question_snapshot=excluded.question_snapshot,

                  state='NEEDS_REVIEW',

                  updated_at=excluded.updated_at

                """, (new_id, user_id, question_id, session["subject_id"], session["grade"],

                      selected_clean, correct_option, explanation, snapshot_json, now, now))



        return self.get_session(user_id, session_id)



    def complete_session(self, user_id, session_id):

        with self.db.transaction() as db:

            session = db.execute("SELECT * FROM practice_sessions WHERE id=? AND user_id=?", (session_id, user_id)).fetchone()

            if not session:

                error("SESSION_NOT_FOUND", 404)

            if session["state"] == "COMPLETED":

                # Idempotent completion: return existing completed session

                return self.get_session(user_id, session_id)

            try:

                snapshots = json.loads(session["question_snapshots"])

            except Exception:

                snapshots = []

            if len(snapshots) == 0:

                raise HTTPException(400, detail={"code": "EMPTY_SESSION", "message": "Phiên làm bài không có câu hỏi."})



            try:

                answers = json.loads(session["answers"])

            except Exception:

                answers = {}



            snapshot_qids = [q["id"] for q in snapshots]

            unanswered_qids = [

                qid for qid in snapshot_qids

                if qid not in answers or not isinstance(answers.get(qid), dict) or answers[qid].get("selectedOption") not in ("A", "B", "C", "D")

            ]

            if unanswered_qids:

                answered_count = len(snapshot_qids) - len(unanswered_qids)

                raise HTTPException(400, detail={

                    "code": "SESSION_INCOMPLETE",

                    "message": f"Bài luyện chưa hoàn thành: mới trả lời {answered_count}/{len(snapshot_qids)} câu hỏi."

                })



            now = time.time()

            db.execute("UPDATE practice_sessions SET state='COMPLETED', updated_at=? WHERE id=?", (now, session_id))

        return self.get_session(user_id, session_id)



    def list_sessions(self, user_id, quiz_id=None, subject_id=None, grade=None, state=None, limit=20):

        query = "SELECT * FROM practice_sessions WHERE user_id=?"

        params = [user_id]

        if quiz_id:

            query += " AND quiz_id=?"

            params.append(quiz_id)

        if subject_id:

            norm_sub = SUBJECT_ALIASES.get(subject_id.strip().lower(), subject_id.strip().lower())

            query += " AND subject_id=?"

            params.append(norm_sub)

        if grade:

            query += " AND grade=?"

            params.append(grade)

        if state and state.upper() != "ALL":

            query += " AND state=?"

            params.append(state.upper().strip())

        query += " ORDER BY updated_at DESC LIMIT ?"

        params.append(limit)



        with self.db.read() as db:

            rows = db.execute(query, params).fetchall()

            items = []

            for r in rows:

                try:

                    snapshots = json.loads(r["question_snapshots"])

                except Exception:

                    snapshots = []

                try:

                    answers = json.loads(r["answers"])

                except Exception:

                    answers = {}



                first_unanswered_idx = 0

                for idx, q in enumerate(snapshots):

                    if q["id"] not in answers:

                        first_unanswered_idx = idx

                        break

                else:

                    first_unanswered_idx = max(0, len(snapshots) - 1)



                items.append({

                    "id": r["id"],

                    "quizId": r["quiz_id"],

                    "subjectId": r["subject_id"],

                    "grade": r["grade"],

                    "state": r["state"],

                    "score": r["score"],

                    "correctCount": r["correct_count"],

                    "totalCount": r["total_count"],

                    "answeredCount": len(answers),

                    "firstUnansweredIndex": first_unanswered_idx,

                    "createdAt": iso(r["created_at"]),

                    "updatedAt": iso(r["updated_at"])

                })

            return {"items": items, "count": len(items)}



    def get_wrong_questions(self, user_id, subject_id=None, grade=None, state=None):

        query = """

        SELECT w.*,

               q.prompt AS legacy_prompt,

               q.option_a AS legacy_a,

               q.option_b AS legacy_b,

               q.option_c AS legacy_c,

               q.option_d AS legacy_d,

               q.chapter AS legacy_chapter,

               q.topic AS legacy_topic,

               q.difficulty AS legacy_difficulty

        FROM wrong_questions_notebook w

        LEFT JOIN questions q ON w.question_id = q.id

        WHERE w.user_id = ?

        """

        params = [user_id]

        if subject_id:

            norm_sub = SUBJECT_ALIASES.get(subject_id.strip().lower(), subject_id.strip().lower())

            query += " AND w.subject_id = ?"

            params.append(norm_sub)

        if grade:

            query += " AND w.grade = ?"

            params.append(grade)

        if state and state.upper() != "ALL":

            query += " AND w.state = ?"

            params.append(state.upper().strip())

        query += " ORDER BY COALESCE(w.updated_at, w.created_at) DESC"



        with self.db.read() as db:

            rows = db.execute(query, params).fetchall()

            items = []

            for r in rows:

                snapshot = None

                if r["question_snapshot"]:

                    try:

                        snapshot = json.loads(r["question_snapshot"])

                    except Exception:

                        snapshot = None



                if snapshot and isinstance(snapshot, dict):

                    has_snapshot = True

                    prompt = snapshot.get("prompt") or "[Nội dung câu hỏi không còn khả dụng]"

                    options = snapshot.get("options") or [

                        {"id": "A", "text": ""},

                        {"id": "B", "text": ""},

                        {"id": "C", "text": ""},

                        {"id": "D", "text": ""}

                    ]

                    # AI snapshots store options by letter; the shared notebook contract
                    # always returns the same array shape consumed by Android.
                    if isinstance(options, dict):
                        options = [{"id": k, "text": str(options.get(k, ""))} for k in ("A", "B", "C", "D")]
                    chapter = snapshot.get("chapter") or "Chưa phân loại"

                    topic = snapshot.get("topic") or "Chủ đề chung"

                    explanation = snapshot.get("explanation") or r["explanation"]

                    difficulty = snapshot.get("difficulty") or "MEDIUM"

                    correct_option = snapshot.get("correctOption") or r["correct_option"]

                else:

                    has_snapshot = False

                    prompt = r["legacy_prompt"] or "[Nội dung câu hỏi không còn khả dụng]"

                    options = [

                        {"id": "A", "text": r["legacy_a"] or ""},

                        {"id": "B", "text": r["legacy_b"] or ""},

                        {"id": "C", "text": r["legacy_c"] or ""},

                        {"id": "D", "text": r["legacy_d"] or ""}

                    ]

                    chapter = r["legacy_chapter"] or "Chưa phân loại"

                    topic = r["legacy_topic"] or "Chủ đề chung"

                    explanation = r["explanation"]

                    difficulty = r["legacy_difficulty"] or "MEDIUM"

                    correct_option = r["correct_option"]



                items.append({

                    "id": r["id"],

                    "questionId": r["question_id"],
                    "source": r["source"] if "source" in r.keys() else "LIBRARY",

                    "subjectId": r["subject_id"],

                    "grade": r["grade"],

                    "chapter": chapter,

                    "topic": topic,

                    "prompt": prompt,

                    "options": options,

                    "selectedOption": r["selected_option"],

                    "correctOption": correct_option,

                    "explanation": explanation,

                    "difficulty": difficulty,

                    "state": r["state"] or "NEEDS_REVIEW",

                    "hasSnapshot": has_snapshot,

                    "canRetry": has_snapshot,

                    "createdAt": iso(r["created_at"]),

                    "updatedAt": iso(r["updated_at"]) if r["updated_at"] else None

                })

            return items



    def retry_wrong_question(self, user_id, entry_id, selected_option):

        if not isinstance(selected_option, str):

            raise HTTPException(400, detail={"code": "INVALID_OPTION", "message": "Phương án chọn phải là chuỗi A, B, C hoặc D."})

        selected_clean = selected_option.strip().upper()

        if selected_clean not in ("A", "B", "C", "D"):

            raise HTTPException(400, detail={"code": "INVALID_OPTION", "message": "Phương án chọn phải là A, B, C hoặc D."})



        with self.db.transaction() as db:

            entry = db.execute("SELECT * FROM wrong_questions_notebook WHERE id=? AND user_id=?", (entry_id, user_id)).fetchone()

            if not entry:

                raise HTTPException(404, detail={"code": "ENTRY_NOT_FOUND", "message": "Không tìm thấy mục sổ lỗi sai."})



            if not entry["question_snapshot"]:

                raise HTTPException(400, detail={

                    "code": "LEGACY_ENTRY_CANNOT_RETRY",

                    "message": "Mục sổ lỗi sai cũ thiếu dữ liệu snapshot nên không thể luyện lại an toàn. Vui lòng làm bài luyện mới để cập nhật sổ lỗi."

                })



            try:

                snapshot = json.loads(entry["question_snapshot"])

            except Exception:

                raise HTTPException(400, detail={

                    "code": "LEGACY_ENTRY_CANNOT_RETRY",

                    "message": "Dữ liệu snapshot của câu hỏi bị lỗi."

                })



            correct_option = snapshot.get("correctOption") or entry["correct_option"]

            if not correct_option:

                raise HTTPException(400, detail={

                    "code": "LEGACY_ENTRY_CANNOT_RETRY",

                    "message": "Dữ liệu đáp án chuẩn của câu hỏi không khả dụng."

                })



            is_correct = (selected_clean == correct_option.strip().upper())

            new_state = "RESOLVED" if is_correct else "NEEDS_REVIEW"

            now = time.time()



            db.execute("""

            UPDATE wrong_questions_notebook

            SET state=?, selected_option=?, updated_at=?

            WHERE id=? AND user_id=?

            """, (new_state, selected_clean, now, entry_id, user_id))



            explanation = snapshot.get("explanation") or entry["explanation"]



            return {

                "entryId": entry["id"],

                "questionId": entry["question_id"],

                "selectedOption": selected_clean,

                "correctOption": correct_option,

                "isCorrect": is_correct,

                "state": new_state,

                "explanation": explanation,

                "updatedAt": iso(now)

            }



    # --- ADMIN FEATURES ---

    def list_questions_admin(self, subject_id=None, grade=None, status=None, difficulty=None, chapter=None, topic=None):

        query = "SELECT * FROM questions WHERE 1=1"

        params = []

        if subject_id:

            norm_sub = SUBJECT_ALIASES.get(subject_id.strip().lower(), subject_id.strip().lower())

            query += " AND subject_id=?"

            params.append(norm_sub)

        if grade is not None:

            query += " AND grade=?"

            params.append(grade)

        if status:

            query += " AND status=?"

            params.append(status.upper())

        if difficulty and difficulty.upper() != "ALL":

            query += " AND difficulty=?"

            params.append(difficulty.upper())

        if chapter:

            query += " AND chapter LIKE ?"

            params.append(f"%{chapter}%")

        if topic:

            query += " AND topic LIKE ?"

            params.append(f"%{topic}%")

        query += " ORDER BY updated_at DESC"



        with self.db.read() as db:

            rows = db.execute(query, params).fetchall()

            items = []

            for r in rows:

                items.append({

                    "id": r["id"],

                    "code": r["code"],

                    "version": r["version"],

                    "subjectId": r["subject_id"],

                    "grade": r["grade"],

                    "chapter": r["chapter"],

                    "topic": r["topic"],

                    "textbook": r["textbook"],

                    "prompt": r["prompt"],

                    "optionA": r["option_a"],

                    "optionB": r["option_b"],

                    "optionC": r["option_c"],

                    "optionD": r["option_d"],

                    "correctOption": r["correct_option"],

                    "explanation": r["explanation"],

                    "difficulty": r["difficulty"],

                    "points": r["points"],

                    "imageUrl": r["image_url"],

                    "source": r["source"],

                    "author": r["author"],

                    "license": r["license"],

                    "status": r["status"],

                    "createdAt": iso(r["created_at"]),

                    "updatedAt": iso(r["updated_at"])

                })

            return items



    def create_or_update_question_admin(self, admin_id, data, expected_version=None):

        now = time.time()

        qid = data.get("id")



        with self.db.transaction() as db:

            existing = None

            if qid:

                existing = db.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()

                if not existing:

                    error("QUESTION_NOT_FOUND", 404)



                # Optimistic concurrency check

                if expected_version is not None and existing["version"] != expected_version:

                    raise HTTPException(409, detail={

                        "code": "VERSION_CONFLICT",

                        "message": f"Xung đột phiên bản: Câu hỏi đã được cập nhật lên v{existing['version']} bởi người khác (phiên bản bạn đang sửa là v{expected_version}). Vui lòng tải lại dữ liệu mới nhất."

                    })



            raw_code = (data.get("code") or "").strip()

            if not raw_code:

                if existing:

                    code = existing["code"]

                else:

                    code = f"Q_{uuid.uuid4().hex[:8].upper()}"

            else:

                code = raw_code



            # Preserve status when editing existing question unless explicitly provided

            existing_status = existing["status"] if existing else "DRAFT"

            new_status = data.get("status") or existing_status



            row_dict = {

                "code": code,

                "subjectId": data.get("subjectId") or (existing["subject_id"] if existing else "physics"),

                "grade": data.get("grade") if data.get("grade") is not None else (existing["grade"] if existing else 10),

                "chapter": data.get("chapter") or (existing["chapter"] if existing else "Tổng hợp"),

                "topic": data.get("topic") or (existing["topic"] if existing else "Chủ đề chung"),

                "textbook": data.get("textbook") or (existing["textbook"] if existing else "Chưa xác định"),

                "prompt": data.get("prompt") if data.get("prompt") is not None else (existing["prompt"] if existing else ""),

                "optionA": data.get("optionA") if data.get("optionA") is not None else (existing["option_a"] if existing else ""),

                "optionB": data.get("optionB") if data.get("optionB") is not None else (existing["option_b"] if existing else ""),

                "optionC": data.get("optionC") if data.get("optionC") is not None else (existing["option_c"] if existing else ""),

                "optionD": data.get("optionD") if data.get("optionD") is not None else (existing["option_d"] if existing else ""),

                "correctOption": data.get("correctOption") if data.get("correctOption") is not None else (existing["correct_option"] if existing else ""),

                "explanation": data.get("explanation") if data.get("explanation") is not None else (existing["explanation"] if existing else ""),

                "difficulty": data.get("difficulty") or (existing["difficulty"] if existing else "MEDIUM"),

                "points": data.get("points") if data.get("points") is not None else (existing["points"] if existing else 10),

                "source": data.get("source") or (existing["source"] if existing else "Admin"),

                "author": data.get("author") or (existing["author"] if existing else "Quản trị viên"),

                "license": data.get("license") or (existing["license"] if existing else "Nội bộ"),

                "status": new_status

            }



            cleaned, err = validate_single_question_dict(row_dict, row_idx=1)

            if err:

                raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": err})



            if existing:

                # Conflict guard: cannot change subject or grade if question belongs to a published quiz

                new_sub = cleaned["subjectId"]

                new_grade = cleaned["grade"]

                old_sub = SUBJECT_ALIASES.get(existing["subject_id"].lower(), existing["subject_id"].lower())

                old_grade = existing["grade"]

                if new_sub != old_sub or new_grade != old_grade:

                    pub_quizzes = db.execute("SELECT id, title, subject_id, grade, question_ids FROM library_quizzes WHERE status='PUBLISHED'").fetchall()

                    for qz in pub_quizzes:

                        try:

                            ids = json.loads(qz["question_ids"])

                        except Exception:

                            ids = []

                        if qid in ids:

                            qz_sub = SUBJECT_ALIASES.get(qz["subject_id"].lower(), qz["subject_id"].lower())

                            if new_sub != qz_sub or new_grade != qz["grade"]:

                                raise HTTPException(409, detail={

                                    "code": "CONFLICT_QUIZ_MEMBERSHIP",

                                    "message": f"Không thể đổi môn/lớp của câu hỏi vì câu hỏi đang thuộc bộ đề đã xuất bản '{qz['title']}' ({qz['subject_id']} lớp {qz['grade']}). Vui lòng gỡ câu hỏi khỏi bộ đề hoặc ẩn bộ đề trước."

                                })



                code_conflict = db.execute("SELECT 1 FROM questions WHERE code=? AND id!=?", (cleaned["code"], qid)).fetchone()

                if code_conflict:

                    raise HTTPException(409, detail={"code": "CODE_CONFLICT", "message": f"Mã câu hỏi '{cleaned['code']}' đã được sử dụng bởi câu hỏi khác"})



                version = existing["version"] + 1

                db.execute("""

                UPDATE questions SET

                    code=?, version=?, subject_id=?, grade=?, chapter=?, topic=?, textbook=?,

                    prompt=?, option_a=?, option_b=?, option_c=?, option_d=?, correct_option=?,

                    explanation=?, difficulty=?, points=?, image_url=?, source=?, author=?, license=?,

                    status=?, updated_at=?

                WHERE id=?

                """, (cleaned["code"], version, cleaned["subjectId"], cleaned["grade"],

                      cleaned["chapter"], cleaned["topic"], cleaned["textbook"],

                      cleaned["prompt"], cleaned["optionA"], cleaned["optionB"],

                      cleaned["optionC"], cleaned["optionD"], cleaned["correctOption"],

                      cleaned["explanation"], cleaned["difficulty"], cleaned["points"],

                      data.get("imageUrl") or existing["image_url"], cleaned["source"], cleaned["author"],

                      cleaned["license"], cleaned["status"], now, qid))

            else:

                code_conflict = db.execute("SELECT 1 FROM questions WHERE code=?", (cleaned["code"],)).fetchone()

                if code_conflict:

                    raise HTTPException(409, detail={"code": "CODE_CONFLICT", "message": f"Mã câu hỏi '{cleaned['code']}' đã tồn tại"})



                new_qid = f"q_{uuid.uuid4().hex[:12]}"

                db.execute("""

                INSERT INTO questions

                (id, code, version, subject_id, grade, chapter, topic, textbook, prompt, option_a, option_b, option_c, option_d, correct_option, explanation, difficulty, points, image_url, source, author, license, status, created_by, created_at, updated_at)

                VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)

                """, (new_qid, cleaned["code"], cleaned["subjectId"], cleaned["grade"],

                      cleaned["chapter"], cleaned["topic"], cleaned["textbook"],

                      cleaned["prompt"], cleaned["optionA"], cleaned["optionB"],

                      cleaned["optionC"], cleaned["optionD"], cleaned["correctOption"],

                      cleaned["explanation"], cleaned["difficulty"], cleaned["points"],

                      data.get("imageUrl"), cleaned["source"], cleaned["author"],

                      cleaned["license"], cleaned["status"], admin_id, now, now))

                qid = new_qid



        return self.get_question_by_id(qid)



    def get_question_by_id(self, qid):

        with self.db.read() as db:

            r = db.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()

            if not r:

                error("QUESTION_NOT_FOUND", 404)

            return {

                "id": r["id"],

                "code": r["code"],

                "version": r["version"],

                "subjectId": r["subject_id"],

                "grade": r["grade"],

                "chapter": r["chapter"],

                "topic": r["topic"],

                "textbook": r["textbook"],

                "prompt": r["prompt"],

                "optionA": r["option_a"],

                "optionB": r["option_b"],

                "optionC": r["option_c"],

                "optionD": r["option_d"],

                "correctOption": r["correct_option"],

                "explanation": r["explanation"],

                "difficulty": r["difficulty"],

                "points": r["points"],

                "imageUrl": r["image_url"],

                "source": r["source"],

                "author": r["author"],

                "license": r["license"],

                "status": r["status"],

                "createdAt": iso(r["created_at"]),

                "updatedAt": iso(r["updated_at"])

            }



    def publish_question_admin(self, qid, publish=True):

        now = time.time()

        with self.db.transaction() as db:

            q = db.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()

            if not q:

                error("QUESTION_NOT_FOUND", 404)



            status = "PUBLISHED" if publish else "DRAFT"

            if publish:

                row_dict = {

                    "code": q["code"],

                    "subjectId": q["subject_id"],

                    "grade": q["grade"],

                    "chapter": q["chapter"],

                    "topic": q["topic"],

                    "textbook": q["textbook"],

                    "prompt": q["prompt"],

                    "optionA": q["option_a"],

                    "optionB": q["option_b"],

                    "optionC": q["option_c"],

                    "optionD": q["option_d"],

                    "correctOption": q["correct_option"],

                    "explanation": q["explanation"],

                    "difficulty": q["difficulty"],

                    "points": q["points"],

                    "source": q["source"],

                    "author": q["author"],

                    "license": q["license"]

                }

                _, err = validate_single_question_dict(row_dict, row_idx=1)

                if err:

                    raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": f"Không thể xuất bản câu hỏi không hợp lệ: {err}"})

            else:

                # When unpublishing a question:

                # Any published quiz containing this question is automatically demoted to DRAFT

                # so that students cannot start an incomplete / truncated quiz!

                pub_quizzes = db.execute("SELECT id, question_ids FROM library_quizzes WHERE status='PUBLISHED'").fetchall()

                for qz in pub_quizzes:

                    try:

                        ids = json.loads(qz["question_ids"])

                    except Exception:

                        ids = []

                    if qid in ids:

                        db.execute("UPDATE library_quizzes SET status='DRAFT' WHERE id=?", (qz["id"],))



            db.execute("UPDATE questions SET status=?, updated_at=? WHERE id=?", (status, now, qid))



        return self.get_question_by_id(qid)



    # --- ADMIN QUIZ MANAGEMENT (PIPELINE TO STUDENT LIBRARY) ---

    def _validate_quiz_payload(self, db, subject_id, grade, q_ids, require_published=False):

        raw_sub = str(subject_id or "physics").strip().lower()

        norm_sub = SUBJECT_ALIASES.get(raw_sub, raw_sub)

        if norm_sub not in ALL_SUBJECTS:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": f"Môn học '{raw_sub}' không hợp lệ"})



        try:

            norm_grade = int(grade)

        except (ValueError, TypeError):

            norm_grade = 10

        if norm_grade not in (10, 11, 12):

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Khối lớp phải là 10, 11 hoặc 12"})



        if not isinstance(q_ids, list) or len(q_ids) == 0:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Bộ đề phải chứa ít nhất 1 câu hỏi"})

        if len(q_ids) > 500:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Bộ đề vượt quá số lượng câu hỏi tối đa (500 câu)"})



        # Check duplicate IDs

        seen_ids = set()

        duplicates = []

        for qid in q_ids:

            if qid in seen_ids:

                duplicates.append(qid)

            seen_ids.add(qid)

        if duplicates:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": f"Danh sách câu hỏi không được chứa ID trùng lặp ({', '.join(duplicates)})"})



        placeholders = ",".join("?" * len(q_ids))

        q_rows = db.execute(f"SELECT id, code, subject_id, grade, status FROM questions WHERE id IN ({placeholders})", q_ids).fetchall()

        found_map = {r["id"]: r for r in q_rows}



        missing = [qid for qid in q_ids if qid not in found_map]

        if missing:

            raise HTTPException(404, detail={"code": "QUESTION_NOT_FOUND", "message": f"Không tìm thấy các câu hỏi: {', '.join(missing)}"})



        # Check subject and grade consistency across all questions

        for qid in q_ids:

            r = found_map[qid]

            r_sub = SUBJECT_ALIASES.get(r["subject_id"].strip().lower(), r["subject_id"].strip().lower())

            if r_sub != norm_sub or r["grade"] != norm_grade:

                raise HTTPException(400, detail={

                    "code": "VALIDATION_FAILED",

                    "message": f"Mã câu hỏi '{r['code'] or r['id']}' thuộc môn {r['subject_id']} lớp {r['grade']}, không khớp với môn {norm_sub} lớp {norm_grade} của bộ đề"

                })



        # Check published status if required

        if require_published:

            draft_qs = [r["code"] or r["id"] for r in q_rows if r["status"] != "PUBLISHED"]

            if draft_qs:

                raise HTTPException(400, detail={

                    "code": "VALIDATION_FAILED",

                    "message": f"Không thể xuất bản bộ đề chứa câu hỏi ở trạng thái DRAFT: {', '.join(draft_qs)}"

                })



        return norm_sub, norm_grade, q_ids, found_map



    def create_or_update_quiz_admin(self, admin_id, data):

        title = str(data.get("title") or "").strip()

        if not title:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Tiêu đề bộ đề không được để trống"})

        if len(title) > 200:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Tiêu đề bộ đề vượt quá 200 ký tự"})



        chapter = str(data.get("chapter") or "Tổng hợp").strip() or "Tổng hợp"

        topic = str(data.get("topic") or "Chủ đề chung").strip() or "Chủ đề chung"

        difficulty = str(data.get("difficulty") or "MEDIUM").strip().upper()

        if difficulty not in ("EASY", "MEDIUM", "HARD"):

            difficulty = "MEDIUM"



        status = str(data.get("status") or "DRAFT").strip().upper()

        if status not in ("DRAFT", "PUBLISHED"):

            status = "DRAFT"



        quiz_id = data.get("id")

        now = time.time()



        with self.db.transaction() as db:

            norm_sub, norm_grade, q_ids, _ = self._validate_quiz_payload(

                db, data.get("subjectId"), data.get("grade"), data.get("questionIds"),

                require_published=(status == "PUBLISHED")

            )



            if quiz_id:

                existing = db.execute("SELECT * FROM library_quizzes WHERE id=?", (quiz_id,)).fetchone()

                if not existing:

                    error("QUIZ_NOT_FOUND", 404)

                db.execute("""

                UPDATE library_quizzes SET

                    title=?, subject_id=?, grade=?, chapter=?, topic=?, difficulty=?, question_ids=?, status=?

                WHERE id=?

                """, (title, norm_sub, norm_grade, chapter, topic, difficulty, json.dumps(q_ids), status, quiz_id))

            else:

                quiz_id = f"quiz_{uuid.uuid4().hex[:12]}"

                db.execute("""

                INSERT INTO library_quizzes (id, title, subject_id, grade, chapter, topic, difficulty, question_ids, status, created_at)

                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)

                """, (quiz_id, title, norm_sub, norm_grade, chapter, topic, difficulty, json.dumps(q_ids), status, now))



        return self.get_quiz_admin_by_id(quiz_id)



    def publish_quiz_admin(self, quiz_id, publish=True):

        with self.db.transaction() as db:

            quiz = db.execute("SELECT * FROM library_quizzes WHERE id=?", (quiz_id,)).fetchone()

            if not quiz:

                error("QUIZ_NOT_FOUND", 404)



            status = "PUBLISHED" if publish else "DRAFT"

            if publish:

                q_ids = json.loads(quiz["question_ids"])

                self._validate_quiz_payload(

                    db, quiz["subject_id"], quiz["grade"], q_ids, require_published=True

                )



            db.execute("UPDATE library_quizzes SET status=? WHERE id=?", (status, quiz_id))



        return self.get_quiz_admin_by_id(quiz_id)



    def list_quizzes_admin(self, subject_id=None, grade=None, status=None):

        query = "SELECT * FROM library_quizzes WHERE 1=1"

        params = []

        if subject_id:

            norm_sub = SUBJECT_ALIASES.get(subject_id.strip().lower(), subject_id.strip().lower())

            query += " AND subject_id=?"

            params.append(norm_sub)

        if grade is not None:

            query += " AND grade=?"

            params.append(grade)

        if status:

            query += " AND status=?"

            params.append(status.upper())

        query += " ORDER BY created_at DESC"



        with self.db.read() as db:

            rows = db.execute(query, params).fetchall()

            quizzes = []

            for r in rows:

                q_ids = json.loads(r["question_ids"])

                quizzes.append({

                    "id": r["id"],

                    "title": r["title"],

                    "subjectId": r["subject_id"],

                    "grade": r["grade"],

                    "chapter": r["chapter"],

                    "topic": r["topic"],

                    "difficulty": r["difficulty"],

                    "questionIds": q_ids,

                    "questionCount": len(q_ids),

                    "status": r["status"],

                    "createdAt": iso(r["created_at"])

                })

            return quizzes



    def get_quiz_admin_by_id(self, quiz_id):

        with self.db.read() as db:

            r = db.execute("SELECT * FROM library_quizzes WHERE id=?", (quiz_id,)).fetchone()

            if not r:

                error("QUIZ_NOT_FOUND", 404)

            q_ids = json.loads(r["question_ids"])

            return {

                "id": r["id"],

                "title": r["title"],

                "subjectId": r["subject_id"],

                "grade": r["grade"],

                "chapter": r["chapter"],

                "topic": r["topic"],

                "difficulty": r["difficulty"],

                "questionIds": q_ids,

                "questionCount": len(q_ids),

                "status": r["status"],

                "createdAt": iso(r["created_at"])

            }



    def generate_excel_template(self, format="xlsx"):

        headers = [

            "Mã câu", "Môn (wireId)", "Lớp (10-12)", "Chương", "Chủ đề", "Bộ sách",

            "Câu hỏi", "Đáp án A", "Đáp án B", "Đáp án C", "Đáp án D", "Đáp án đúng (A/B/C/D)",

            "Lời giải", "Độ khó (EASY/MEDIUM/HARD)", "Điểm", "Nguồn", "Tác giả", "Quyền"

        ]

        sample_row = [

            "PHY10_01", "physics", 10, "Chương 1: Động học", "Chuyển động thẳng đều", "Kết nối tri thức",

            "Phương trình chuyển động x = 5 + 6t (m;s). Vận tốc là:",

            "5 m/s", "6 m/s", "-5 m/s", "11 m/s", "B",

            "So sánh với x = x0 + vt có v = 6 m/s", "EASY", 10, "Bộ đề mẫu", "Giáo viên biên soạn", "Nội bộ"

        ]



        if format == "csv":

            buf = io.StringIO()

            writer = csv.writer(buf)

            writer.writerow(headers)

            writer.writerow(sample_row)

            return ("\ufeff" + buf.getvalue()).encode("utf-8")



        wb = openpyxl.Workbook()

        ws = wb.active

        ws.title = "Câu hỏi"

        ws.append(headers)

        ws.append(sample_row)

        buf = io.BytesIO()

        wb.save(buf)

        return buf.getvalue()



    def parse_import_file(self, content_bytes):

        if not content_bytes or len(content_bytes.strip()) == 0:

            return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "File rỗng hoặc không có dữ liệu"}]}



        raw_rows = []

        formula_errors = []



        is_xlsx = content_bytes.startswith(b"PK\x03\x04")

        if is_xlsx:

            try:

                with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:

                    total_uncompressed = sum(info.file_size for info in zf.infolist())

                    if total_uncompressed > 25 * 1024 * 1024:

                        return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "File giải nén vượt quá giới hạn an toàn (tối đa 25MB)"}]}

                    compressed_size = sum(info.compress_size for info in zf.infolist())

                    if compressed_size > 0 and (total_uncompressed / compressed_size) > 50:

                        return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "Tỷ lệ nén file bất thường, bị từ chối để đảm bảo an toàn"}]}

            except zipfile.BadZipFile:

                return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "File .xlsx bị hỏng hoặc không đúng định dạng zip"}]}



            try:

                wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=False, read_only=True)

                candidates = [s for s in wb.worksheets if s.title in ("Câu hỏi", "Questions", "Sheet1")]

                ws = candidates[0] if candidates else (wb.worksheets[0] if wb.worksheets else None)

                if ws is None:

                    wb.close()

                    return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "File Excel không có trang tính (sheet) nào"}]}



                for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):

                    if r_idx > 2001:

                        wb.close()

                        return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": r_idx, "message": "File Excel vượt quá số lượng câu hỏi tối đa cho phép (tối đa 2000 câu hỏi)"}]}



                    row_vals = []

                    for c_idx, cell in enumerate(row, start=1):

                        val = cell.value

                        if cell.data_type == 'f' or (isinstance(val, str) and val.strip().startswith("=")):

                            formula_errors.append({"row": r_idx, "message": f"Dòng {r_idx}, Cột {c_idx}: Phát hiện công thức/macro Excel không được phép ('{val}')"})

                        val_str = "" if val is None else str(val).strip()

                        row_vals.append(val_str)

                    raw_rows.append(row_vals)

                wb.close()

            except Exception as e:

                return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": f"Không thể đọc file Excel: {str(e)}"}]}

        else:

            if b"\x00" in content_bytes[:4096]:

                return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "Định dạng file không được hỗ trợ hoặc file nhị phân không hợp lệ. Vui lòng sử dụng file .csv hoặc .xlsx."}]}



            try:

                text = content_bytes.decode("utf-8-sig")

            except UnicodeDecodeError:

                return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "File không đúng định dạng mã hoá UTF-8"}]}



            try:

                stream = io.StringIO(text)

                reader = csv.reader(stream)

                for r_idx, row in enumerate(reader, start=1):

                    if r_idx > 2001:

                        return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": r_idx, "message": "File CSV vượt quá số lượng câu hỏi tối đa cho phép (tối đa 2000 câu hỏi)"}]}

                    raw_rows.append([col.strip() for col in row])

            except csv.Error as e:

                return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": f"Lỗi cú pháp file CSV: {str(e)}"}]}



        if formula_errors:

            return {"totalRows": len(raw_rows) - 1 if len(raw_rows) > 1 else 0, "validCount": 0, "errorCount": len(formula_errors), "validRows": [], "errors": formula_errors}



        if not raw_rows or not any(any(c for c in r) for r in raw_rows):

            return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 0, "message": "File rỗng hoặc không có dữ liệu"}]}



        header_row = raw_rows[0]

        mapped_headers = {}

        seen_col_names = {}

        header_errors = []



        for c_idx, raw_name in enumerate(header_row):

            norm_name = raw_name.strip().lower()

            if not norm_name:

                continue

            canonical = COLUMN_MAP.get(norm_name)

            if canonical:

                if canonical in mapped_headers:

                    header_errors.append(f"Tiêu đề cột '{raw_name}' bị trùng lặp (cột {c_idx + 1} và cột {seen_col_names[canonical] + 1})")

                else:

                    mapped_headers[canonical] = c_idx

                    seen_col_names[canonical] = c_idx



        missing_required = [label for key, label in REQUIRED_FIELD_LABELS.items() if key not in mapped_headers]

        if missing_required:

            is_probably_header = ("prompt" in mapped_headers and sum(1 for k in ("optionA", "optionB", "optionC", "optionD") if k in mapped_headers) >= 2)

            if not is_probably_header:

                return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 1, "message": "Dòng 1: File thiếu dòng tiêu đề (header) hợp lệ. Vui lòng tải file mẫu chuẩn từ hệ thống."}]}

            return {"totalRows": 0, "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 1, "message": f"Dòng 1: File thiếu các cột bắt buộc: {', '.join(missing_required)}"}]}



        if header_errors:

            return {"totalRows": 0, "validCount": 0, "errorCount": len(header_errors), "validRows": [], "errors": [{"row": 1, "message": f"Dòng 1: {err}"} for err in header_errors]}



        data_rows = [r for r in raw_rows[1:] if any(c for c in r)]

        if not data_rows:

            return {"totalRows": 0, "validCount": 0, "errorCount": 0, "validRows": [], "errors": [{"row": 1, "message": "File chỉ có dòng tiêu đề, chưa có dữ liệu câu hỏi nào"}]}

        if len(data_rows) > 2000:

            return {"totalRows": len(data_rows), "validCount": 0, "errorCount": 1, "validRows": [], "errors": [{"row": 2002, "message": "Số lượng câu hỏi trong file vượt quá giới hạn tối đa (tối đa 2000 câu hỏi)"}]}



        with self.db.read() as db:

            existing_codes = {r[0] for r in db.execute("SELECT code FROM questions WHERE code IS NOT NULL").fetchall()}



        valid_rows = []

        errors = []

        seen_codes = set()



        for idx, row in enumerate(raw_rows[1:], start=2):

            if not row or not any(str(c).strip() for c in row):

                continue



            row_dict = {}

            for field, col_idx in mapped_headers.items():

                if col_idx < len(row):

                    row_dict[field] = row[col_idx]

                else:

                    row_dict[field] = ""



            cleaned, err = validate_single_question_dict(row_dict, row_idx=idx, seen_codes=seen_codes, existing_db_codes=existing_codes)

            if err:

                errors.append({"row": idx, "message": err})

            else:

                valid_rows.append(cleaned)



        return {

            "totalRows": len(data_rows),

            "validCount": len(valid_rows),

            "errorCount": len(errors),

            "validRows": valid_rows,

            "errors": errors

        }



    def confirm_import(self, admin_id, valid_rows, publish=False, idempotency_key=None):

        if not isinstance(valid_rows, list) or len(valid_rows) == 0:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Danh sách câu hỏi cần import không được rỗng"})



        if len(valid_rows) > 2000:

            raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": "Số lượng câu hỏi import vượt quá giới hạn tối đa (2000 câu)"})



        payload_repr = {"rows": valid_rows, "publish": bool(publish)}

        request_hash = hashlib.sha256(dump(payload_repr).encode("utf-8")).hexdigest()



        raw_key = (idempotency_key or "").strip()

        if raw_key:

            if len(raw_key) > 128:

                raise HTTPException(400, detail={"code": "INVALID_IDEMPOTENCY_KEY", "message": "Idempotency-Key vượt quá độ dài tối đa (128 ký tự)"})

            effective_key = raw_key

        else:

            effective_key = f"legacy_compat_{request_hash[:32]}"



        with self.db.transaction() as db:

            cached = db.execute(

                "SELECT request_hash, response FROM idempotency WHERE owner=? AND operation='ADMIN_IMPORT' AND key=?",

                (admin_id, effective_key)

            ).fetchone()

            if cached:

                if cached["request_hash"] != request_hash:

                    raise HTTPException(409, detail={

                        "code": "IDEMPOTENCY_CONFLICT",

                        "message": "Định danh thao tác import đã được sử dụng với payload hoặc tùy chọn xuất bản khác."

                    })

                return json.loads(cached["response"])



            seen_codes = set()

            cleaned_batch = []

            for idx, item in enumerate(valid_rows, start=1):

                cleaned, err = validate_single_question_dict(item, row_idx=idx, seen_codes=seen_codes, existing_db_codes=None)

                if err:

                    raise HTTPException(400, detail={"code": "VALIDATION_FAILED", "message": err, "row": idx})

                cleaned_batch.append(cleaned)



            status = "PUBLISHED" if bool(publish) else "DRAFT"

            now = time.time()

            imported = []



            all_codes = [q["code"] for q in cleaned_batch]

            placeholders = ",".join("?" * len(all_codes))

            existing_in_db = db.execute(f"SELECT code FROM questions WHERE code IN ({placeholders})", all_codes).fetchall()

            if existing_in_db:

                conflict_codes = [r["code"] for r in existing_in_db]

                raise HTTPException(409, detail={"code": "CODE_CONFLICT", "message": f"Mã câu hỏi đã tồn tại trong hệ thống: {', '.join(conflict_codes)}. Không thể import đè."})



            for item in cleaned_batch:

                qid = f"q_{uuid.uuid4().hex[:12]}"

                try:

                    db.execute("""

                    INSERT INTO questions

                    (id, code, version, subject_id, grade, chapter, topic, textbook, prompt, option_a, option_b, option_c, option_d, correct_option, explanation, difficulty, points, image_url, source, author, license, status, created_by, created_at, updated_at)

                    VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, ?, ?, ?, ?, ?, ?)

                    """, (qid, item["code"], item["subjectId"], item["grade"], item["chapter"],

                          item["topic"], item["textbook"], item["prompt"], item["optionA"],

                          item["optionB"], item["optionC"], item["optionD"], item["correctOption"],

                          item["explanation"], item["difficulty"], item["points"], item["source"],

                          item["author"], item["license"], status, admin_id, now, now))

                except sqlite3.IntegrityError as e:

                    raise HTTPException(409, detail={"code": "INTEGRITY_CONFLICT", "message": f"Xung đột ràng buộc dữ liệu: {str(e)}"})



                imported.append({

                    "id": qid,

                    "code": item["code"],

                    "version": 1,

                    "subjectId": item["subjectId"],

                    "grade": item["grade"],

                    "prompt": item["prompt"],

                    "correctOption": item["correctOption"],

                    "status": status,

                    "createdAt": iso(now),

                    "updatedAt": iso(now)

                })



            result = {"importedCount": len(imported), "questions": imported}

            db.execute(

                "INSERT INTO idempotency (owner, operation, key, request_hash, response) VALUES (?, 'ADMIN_IMPORT', ?, ?, ?)",

                (admin_id, effective_key, request_hash, dump(result))

            )

            return result
