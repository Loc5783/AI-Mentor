import io
import pytest
import sqlite3
import openpyxl
from fastapi.testclient import TestClient
from aimentor.config import Settings
from aimentor.main import create_app
from aimentor.manage_roles import main as manage_roles_main
from test_api import signup, headers

def setup_client(tmp_path):
    app = create_app(Settings(tmp_path))
    client = TestClient(app)

    # 1. Student user
    student = signup(client, "StudentUser")
    student_h = headers(student)

    # 2. Admin user
    admin_u = signup(client, "AdminUser")
    admin_id = admin_u["userId"]
    db_file = tmp_path / "aimentor.sqlite3"

    # Promote to ADMIN using manage_roles
    import sys
    old_argv = sys.argv
    sys.argv = ["manage_roles", "--db-path", str(db_file), "--user", admin_id, "--role", "ADMIN", "--apply"]
    try:
        manage_roles_main()
    except SystemExit:
        pass
    finally:
        sys.argv = old_argv

    login_res = client.post("/auth/sign-in", json={"identifier": "adminuser@school.example", "password": "Mentor01!2026"})
    admin_token = login_res.json()["accessToken"]
    admin_h = {"Authorization": f"Bearer {admin_token}"}

    return client, admin_h, student_h, db_file

# =========================================================================
# GROUP A: VALID FILES & TEMPLATES
# =========================================================================

def test_valid_csv_utf8_with_vietnamese_and_multiline(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    csv_content = (
        'Mã câu,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,Đáp án A,Đáp án B,Đáp án C,Đáp án D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n'
        'PHY_VN_01,physics,10,Chương 1: Mở đầu,Động học,Kết nối tri thức,"Một chất điểm chuyển động có phương trình:\nx = 10 + 5t (m; s).\nVận tốc của vật là:",'
        '"5 m/s","10 m/s","15 m/s","-5 m/s",A,"Lời giải chi tiết:\nĐối chiếu x = x0 + vt có v = 5 m/s.",EASY,10,Đề thi thử,Thầy Nguyễn Văn A,Bản quyền mở\n'
    )
    # Preview
    res = client.post("/admin/import/preview", headers=admin_h, content=csv_content.encode("utf-8"))
    assert res.status_code == 200
    data = res.json()
    assert data["totalRows"] == 1
    assert data["validCount"] == 1
    assert data["errorCount"] == 0

    # Ensure multiline newline is preserved!
    q = data["validRows"][0]
    assert "\n" in q["prompt"]
    assert "x = 10 + 5t" in q["prompt"]
    assert "\n" in q["explanation"]
    assert q["author"] == "Thầy Nguyễn Văn A"

    # Confirm
    res_conf = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": data["validRows"], "publish": False})
    assert res_conf.status_code == 200
    assert res_conf.json()["importedCount"] == 1

    # Verify in DB
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT prompt, explanation, author, status FROM questions WHERE code='PHY_VN_01'")
    row = cur.fetchone()
    conn.close()
    assert row is not None
    assert "\n" in row[0]
    assert "\n" in row[1]
    assert row[2] == "Thầy Nguyễn Văn A"
    assert row[3] == "DRAFT"

def test_valid_csv_utf8_bom(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)
    csv_content = (
        'Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n'
        'Q_BOM_01,chemistry,10,Chương 1,Nguyên tử,KNTT,Số proton trong hạt nhân nhôm là:,13,14,27,10,A,Al có Z=13,EASY,10,SBT,Tác giả,Nội bộ\n'
    )
    bom_bytes = b"\xef\xbb\xbf" + csv_content.encode("utf-8")
    res = client.post("/admin/import/preview", headers=admin_h, content=bom_bytes)
    assert res.status_code == 200
    assert res.json()["validCount"] == 1
    assert res.json()["errorCount"] == 0

def test_template_download_and_roundtrip(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # 1. Download XLSX template
    res_xlsx = client.get("/admin/template/excel?format=xlsx")
    assert res_xlsx.status_code == 200
    assert "openxmlformats" in res_xlsx.headers["content-type"]
    assert "Mau_Cau_Hoi_AI_Mentor.xlsx" in res_xlsx.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert "Câu hỏi" in wb.sheetnames
    wb.close()

    # 2. Download CSV template
    res_csv = client.get("/admin/template/excel?format=csv")
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    content_csv = res_csv.content.decode("utf-8-sig")
    assert "Mã câu" in content_csv
    assert "PHY10_01" in content_csv

    # 3. Add a new row to CSV template and import
    custom_csv = content_csv + "BIO10_NEW,biology,10,Chương 1,Tế bào,KNTT,Thành phần cấu tạo nên màng sinh chất là:,Phospholipid và protein,Chỉ protein,Chỉ lipid,ADN,A,Màng sinh chất cấu tạo từ lớp kép phospholipid và protein,MEDIUM,10,Nguồn,Tác giả,Nội bộ\n"
    res_prev = client.post("/admin/import/preview", headers=admin_h, content=custom_csv.encode("utf-8"))
    assert res_prev.status_code == 200
    prev_data = res_prev.json()
    assert prev_data["validCount"] >= 1

    # Confirm
    res_conf = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": prev_data["validRows"], "publish": True})
    assert res_conf.status_code == 200
    assert res_conf.json()["importedCount"] == prev_data["validCount"]

def test_valid_xlsx_openpyxl_import(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Câu hỏi"
    ws.append(["Mã câu", "Môn", "Lớp", "Chương", "Chủ đề", "Bộ sách", "Câu hỏi", "A", "B", "C", "D", "Đáp án đúng", "Lời giải", "Độ khó", "Điểm", "Nguồn", "Tác giả", "Quyền"])
    ws.append(["XLSX_01", "physics", 10, "Chương 1", "Chuyển động", "KNTT", "Đơn vị vận tốc trong hệ SI là:\nm/s hay km/h?", "m/s", "km/h", "cm/s", "m/min", "A", "Hệ SI dùng m/s", "EASY", 10, "SGK", "Bộ GD", "Công cộng"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    res = client.post("/admin/import/preview", headers=admin_h, content=xlsx_bytes)
    assert res.status_code == 200
    data = res.json()
    assert data["validCount"] == 1
    assert data["errorCount"] == 0
    assert "\n" in data["validRows"][0]["prompt"]

    # Confirm
    res_conf = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": data["validRows"], "publish": True})
    assert res_conf.status_code == 200
    assert res_conf.json()["importedCount"] == 1

# =========================================================================
# GROUP B: VALIDATION & ERROR REPORTING
# =========================================================================

def test_missing_header_row_rejected(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # File without header, directly having a data row
    no_header_csv = "Q01,physics,10,Chương 1,Topic,KNTT,Câu hỏi không có header?,A1,A2,A3,A4,A,Expl,EASY,10,Src,Auth,Lic\n"
    res = client.post("/admin/import/preview", headers=admin_h, content=no_header_csv.encode("utf-8"))
    assert res.status_code == 200
    data = res.json()
    assert data["validCount"] == 0
    assert data["errorCount"] >= 1
    assert "File thiếu dòng tiêu đề" in data["errors"][0]["message"]

def test_missing_required_columns_in_header(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # Header missing "Đáp án đúng"
    bad_header = "Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n"
    res = client.post("/admin/import/preview", headers=admin_h, content=bad_header.encode("utf-8"))
    assert res.status_code == 200
    data = res.json()
    assert data["errorCount"] >= 1
    assert "Đáp án đúng" in data["errors"][0]["message"]

def test_duplicate_header_columns(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # Header has two "Câu hỏi" columns
    dup_header = "Mã,Câu hỏi,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải\n"
    res = client.post("/admin/import/preview", headers=admin_h, content=dup_header.encode("utf-8"))
    assert res.status_code == 200
    data = res.json()
    assert any("trùng lặp" in e["message"] for e in data["errors"])

def test_empty_file_and_header_only_file(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # 1. 0 bytes
    res_empty = client.post("/admin/import/preview", headers=admin_h, content=b"")
    assert res_empty.status_code == 200
    assert res_empty.json()["errorCount"] == 1
    assert "File rỗng" in res_empty.json()["errors"][0]["message"]

    # 2. Header only
    hdr_only = "Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n"
    res_hdr = client.post("/admin/import/preview", headers=admin_h, content=hdr_only.encode("utf-8"))
    assert res_hdr.status_code == 200
    data_hdr = res_hdr.json()
    assert data_hdr["validCount"] == 0
    assert any("chưa có dữ liệu" in e["message"] for e in data_hdr["errors"])

def test_row_validation_errors(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    csv_data = (
        'Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n'
        'Q_ERR_01,physics,10,Ch1,T1,KNTT,,A1,B1,C1,D1,A,Expl,EASY,10,Src,Auth,Lic\n'               # Empty prompt
        'Q_ERR_02,physics,10,Ch1,T1,KNTT,Prompt 2?,A1,,C1,D1,A,Expl,EASY,10,Src,Auth,Lic\n'          # Missing Option B
        'Q_ERR_03,physics,10,Ch1,T1,KNTT,Prompt 3?,A1,B1,C1,D1,E,Expl,EASY,10,Src,Auth,Lic\n'       # Invalid correctOption 'E'
        'Q_ERR_04,unknown_sub,10,Ch1,T1,KNTT,Prompt 4?,A1,B1,C1,D1,A,Expl,EASY,10,Src,Auth,Lic\n'   # Invalid subject
        'Q_ERR_05,physics,99,Ch1,T1,KNTT,Prompt 5?,A1,B1,C1,D1,A,Expl,EASY,10,Src,Auth,Lic\n'       # Invalid grade 99
        'Q_ERR_06,physics,10,Ch1,T1,KNTT,Prompt 6?,A1,B1,C1,D1,A,Expl,SUPER_HARD,10,Src,Auth,Lic\n' # Invalid difficulty
        'Q_ERR_07,physics,10,Ch1,T1,KNTT,Prompt 7?,A1,B1,C1,D1,A,Expl,EASY,-5,Src,Auth,Lic\n'        # Negative points
    )
    res = client.post("/admin/import/preview", headers=admin_h, content=csv_data.encode("utf-8"))
    assert res.status_code == 200
    data = res.json()
    assert data["validCount"] == 0
    assert data["errorCount"] == 7

def test_xlsx_formula_cells_rejected(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Mã", "Môn", "Lớp", "Chương", "Chủ đề", "Bộ sách", "Câu hỏi", "A", "B", "C", "D", "Đáp án đúng", "Lời giải", "Độ khó", "Điểm", "Nguồn", "Tác giả", "Quyền"])
    # Put formula in prompt
    ws.append(["FORMULA_01", "physics", 10, "Ch1", "T1", "KNTT", "=SUM(A1:B10)", "A", "B", "C", "D", "A", "Expl", "EASY", 10, "Src", "Auth", "Lic"])
    buf = io.BytesIO()
    wb.save(buf)

    res = client.post("/admin/import/preview", headers=admin_h, content=buf.getvalue())
    assert res.status_code == 200
    data = res.json()
    assert data["validCount"] == 0
    assert any("công thức/macro" in e["message"] for e in data["errors"])

# =========================================================================
# GROUP C: ATOMICITY, ROLLBACK & DUPLICATE PROTECTION
# =========================================================================

def test_confirm_validates_and_rejects_tampered_payload(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # Directly calling confirm with invalid row (e.g. empty prompt, invalid correctOption)
    tampered_body = {
        "validRows": [
            {
                "code": "HACK_01",
                "subjectId": "physics",
                "grade": 10,
                "prompt": "",  # Empty prompt!
                "optionA": "Opt1",
                "optionB": "Opt2",
                "optionC": "Opt3",
                "optionD": "Opt4",
                "correctOption": "INVALID",
                "points": 9999
            }
        ],
        "publish": True
    }
    res = client.post("/admin/import/confirm", headers=admin_h, json=tampered_body)
    assert res.status_code in (400, 422)

    # Ensure NOT inserted into database
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM questions WHERE code='HACK_01'")
    assert cur.fetchone()[0] == 0
    conn.close()

def test_batch_atomicity_all_or_nothing_rollback(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # Batch with 1 valid question and 1 invalid question
    batch = {
        "validRows": [
            {
                "code": "VALID_ROW_01",
                "subjectId": "physics",
                "grade": 10,
                "prompt": "Câu hỏi số 1 hợp lệ?",
                "optionA": "1", "optionB": "2", "optionC": "3", "optionD": "4",
                "correctOption": "A",
                "points": 10
            },
            {
                "code": "INVALID_ROW_02",
                "subjectId": "physics",
                "grade": 10,
                "prompt": "Câu hỏi số 2 bị lỗi đáp án?",
                "optionA": "1", "optionB": "2", "optionC": "3", "optionD": "4",
                "correctOption": "Z",  # INVALID OPTION!
                "points": 10
            }
        ],
        "publish": False
    }

    res = client.post("/admin/import/confirm", headers=admin_h, json=batch)
    assert res.status_code == 400

    # ATOMICITY PROOF: Ensure Question 1 was NOT committed!
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM questions WHERE code='VALID_ROW_01'")
    assert cur.fetchone()[0] == 0
    conn.close()

def test_duplicate_code_in_same_file_detected(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    csv_data = (
        'Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n'
        'DUP_01,physics,10,Ch1,T1,KNTT,Câu hỏi 1 cùng mã?,A1,B1,C1,D1,A,Expl,EASY,10,Src,Auth,Lic\n'
        'DUP_01,physics,10,Ch1,T1,KNTT,Câu hỏi 2 cùng mã?,A1,B1,C1,D1,B,Expl,EASY,10,Src,Auth,Lic\n'
    )
    res = client.post("/admin/import/preview", headers=admin_h, content=csv_data.encode("utf-8"))
    assert res.status_code == 200
    data = res.json()
    assert any("trùng lặp" in e["message"] and "DUP_01" in e["message"] for e in data["errors"])

def test_duplicate_code_with_existing_db_question_fails_and_preserves_old(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # 1. First import succeeds
    csv_1 = (
        'Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n'
        'EXISTING_CODE_01,physics,10,Ch1,T1,KNTT,Câu hỏi gốc ban đầu?,A1,B1,C1,D1,A,Expl,EASY,10,Src,Auth,Lic\n'
    )
    res1_prev = client.post("/admin/import/preview", headers=admin_h, content=csv_1.encode("utf-8"))
    res1_conf = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": res1_prev.json()["validRows"], "publish": True})
    assert res1_conf.status_code == 200

    # Get original ID and status
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT id, prompt, status, version FROM questions WHERE code='EXISTING_CODE_01'")
    orig = cur.fetchone()
    conn.close()
    assert orig is not None
    orig_id, orig_prompt, orig_status, orig_ver = orig

    # 2. Try importing same code again with different prompt
    csv_2 = (
        'Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n'
        'EXISTING_CODE_01,physics,10,Ch1,T1,KNTT,Câu hỏi mới cố tình đè mã cũ?,A1,B1,C1,D1,B,Expl,EASY,10,Src,Auth,Lic\n'
    )
    res2_prev = client.post("/admin/import/preview", headers=admin_h, content=csv_2.encode("utf-8"))
    assert res2_prev.status_code == 200
    assert any("đã tồn tại" in e["message"] and "EXISTING_CODE_01" in e["message"] for e in res2_prev.json()["errors"])

    # 3. Direct confirm attempt with conflicting code must be rejected with 409 Conflict
    conflicting_batch = [
        {
            "code": "EXISTING_CODE_01",
            "subjectId": "physics",
            "grade": 10,
            "prompt": "Câu hỏi mới cố tình đè mã cũ?",
            "optionA": "A1", "optionB": "B1", "optionC": "C1", "optionD": "D1",
            "correctOption": "B",
            "points": 10
        }
    ]
    res2_conf = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": conflicting_batch, "publish": False})
    assert res2_conf.status_code == 409
    assert res2_conf.json()["code"] == "CODE_CONFLICT"
    assert "đã tồn tại" in res2_conf.json().get("message", "")

    # 4. Verify original question is UNTOUCHED (ID, prompt, status, version intact!)
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT id, prompt, status, version FROM questions WHERE code='EXISTING_CODE_01'")
    after = cur.fetchone()
    conn.close()
    assert after == (orig_id, orig_prompt, orig_status, orig_ver)

# =========================================================================
# GROUP D: SECURITY & ROLE ACCESS
# =========================================================================

def test_student_and_unauthenticated_access_denied(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    csv_data = "Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải\n"

    # Student cannot preview or confirm (403)
    assert client.post("/admin/import/preview", headers=student_h, content=csv_data.encode()).status_code == 403
    assert client.post("/admin/import/confirm", headers=student_h, json={"validRows": []}).status_code == 403

    # Unauthenticated cannot preview or confirm (401)
    assert client.post("/admin/import/preview", content=csv_data.encode()).status_code == 401
    assert client.post("/admin/import/confirm", json={"validRows": []}).status_code == 401


# =========================================================================
# GROUP E: STRICT FIELD LENGTH LIMITS (NO SILENT TRUNCATION) & REAL ROLLBACK
# =========================================================================

def test_field_length_exact_limit_and_overflow_rejected(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # 1. Exact boundary: 5000 chars explanation, 200 chars source, 200 chars chapter, 200 chars topic, 100 chars textbook, 100 chars author, 100 chars license
    exact_expl = "E" * 5000
    exact_source = "S" * 200
    exact_chapter = "C" * 200
    exact_topic = "T" * 200
    exact_textbook = "B" * 100
    exact_author = "A" * 100
    exact_license = "L" * 100

    row_exact = {
        "code": "EXACT_LEN_01",
        "subjectId": "physics",
        "grade": 10,
        "chapter": exact_chapter,
        "topic": exact_topic,
        "textbook": exact_textbook,
        "prompt": "Câu hỏi kiểm tra giới hạn chuẩn xác",
        "optionA": "Opt A", "optionB": "Opt B", "optionC": "Opt C", "optionD": "Opt D",
        "correctOption": "A",
        "explanation": exact_expl,
        "difficulty": "EASY",
        "points": 10,
        "source": exact_source,
        "author": exact_author,
        "license": exact_license
    }

    # Confirm exact length
    res_exact = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": [row_exact], "publish": False})
    assert res_exact.status_code == 200
    assert res_exact.json()["importedCount"] == 1

    # Verify database preserves exact lengths WITHOUT silent truncation
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT explanation, source, chapter, topic, textbook, author, license FROM questions WHERE code='EXACT_LEN_01'")
    r = cur.fetchone()
    conn.close()
    assert r is not None
    assert len(r[0]) == 5000, f"Expected 5000 chars explanation, got {len(r[0])}"
    assert len(r[1]) == 200, f"Expected 200 chars source, got {len(r[1])}"
    assert len(r[2]) == 200, f"Expected 200 chars chapter, got {len(r[2])}"
    assert len(r[3]) == 200, f"Expected 200 chars topic, got {len(r[3])}"
    assert len(r[4]) == 100, f"Expected 100 chars textbook, got {len(r[4])}"
    assert len(r[5]) == 100, f"Expected 100 chars author, got {len(r[5])}"
    assert len(r[6]) == 100, f"Expected 100 chars license, got {len(r[6])}"

    # 2. Boundary overflow: Limit + 1 character
    overflow_fields = [
        ("explanation", "E" * 5001, "Lời giải (explanation)"),
        ("source", "S" * 201, "Nguồn trích dẫn (source)"),
        ("chapter", "C" * 201, "Tên chương (chapter)"),
        ("topic", "T" * 201, "Chủ đề (topic)"),
        ("textbook", "B" * 101, "Bộ sách (textbook)"),
        ("author", "A" * 101, "Tác giả (author)"),
        ("license", "L" * 101, "Bản quyền (license)")
    ]

    for field_name, overflow_val, label_expected in overflow_fields:
        bad_row = dict(row_exact)
        bad_row["code"] = f"OVER_{field_name.upper()}"
        bad_row[field_name] = overflow_val

        # Preview must reject
        # Create CSV with overflow field
        csv_line = (
            f"Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng,Lời giải,Độ khó,Điểm,Nguồn,Tác giả,Quyền\n"
            f"{bad_row['code']},physics,10,{bad_row['chapter']},{bad_row['topic']},{bad_row['textbook']},Câu hỏi?,A,B,C,D,A,{bad_row['explanation']},EASY,10,{bad_row['source']},{bad_row['author']},{bad_row['license']}\n"
        )
        prev_res = client.post("/admin/import/preview", headers=admin_h, content=csv_line.encode("utf-8"))
        assert prev_res.status_code == 200
        prev_data = prev_res.json()
        assert prev_data["validCount"] == 0
        assert any(label_expected in e["message"] and "vượt quá độ dài tối đa" in e["message"] for e in prev_data["errors"])

        # Confirm directly must reject with 400
        conf_res = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": [bad_row], "publish": False})
        assert conf_res.status_code == 400
        conf_data = conf_res.json()
        assert label_expected in conf_data.get("message", "")
        assert "vượt quá độ dài tối đa" in conf_data.get("message", "")

    # 3. Batch atomicity with overflow: 1 valid + 1 overflow question
    batch_mixed = [
        {
            "code": "BATCH_OK_BEFORE_OVERFLOW",
            "subjectId": "physics",
            "grade": 10,
            "prompt": "Câu hỏi hợp lệ trước câu lỗi độ dài",
            "optionA": "A", "optionB": "B", "optionC": "C", "optionD": "D",
            "correctOption": "A",
            "explanation": "Lời giải ngắn hợp lệ"
        },
        {
            "code": "BATCH_OVERFLOW_EXPLANATION",
            "subjectId": "physics",
            "grade": 10,
            "prompt": "Câu hỏi chứa lời giải vượt giới hạn",
            "optionA": "A", "optionB": "B", "optionC": "C", "optionD": "D",
            "correctOption": "B",
            "explanation": "X" * 5001
        }
    ]
    res_batch_over = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": batch_mixed, "publish": False})
    assert res_batch_over.status_code == 400
    assert "Lời giải (explanation) vượt quá độ dài tối đa" in res_batch_over.json().get("message", "")

    # Prove database wrote NOTHING from this batch
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM questions WHERE code IN ('BATCH_OK_BEFORE_OVERFLOW', 'BATCH_OVERFLOW_EXPLANATION')")
    assert cur.fetchone()[0] == 0, "Database must not write ANY questions when batch validation fails"
    conn.close()

def test_forced_mid_write_rollback_during_transaction(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # Inject an SQLite trigger that simulates an unexpected failure (e.g. disk corruption/abort)
    # on the SECOND question of the batch:
    conn = sqlite3.connect(db_file)
    conn.execute("""
    CREATE TRIGGER simulate_mid_write_failure
    BEFORE INSERT ON questions
    WHEN NEW.code = 'CRASH_ON_SECOND_ROW'
    BEGIN
        SELECT RAISE(ABORT, 'SIMULATED_DISK_WRITE_ABORT_ON_ROW_2');
    END;
    """)
    conn.commit()
    conn.close()

    batch_to_confirm = [
        {
            "code": "COMMITTED_OR_ROLLED_BACK_01",
            "subjectId": "physics",
            "grade": 10,
            "prompt": "Câu hỏi số 1 hợp lệ sẽ được ghi trước khi crash?",
            "optionA": "A1", "optionB": "B1", "optionC": "C1", "optionD": "D1",
            "correctOption": "A",
            "points": 10
        },
        {
            "code": "CRASH_ON_SECOND_ROW",
            "subjectId": "physics",
            "grade": 10,
            "prompt": "Câu hỏi số 2 gây crash giữa chừng lúc ghi DB",
            "optionA": "A2", "optionB": "B2", "optionC": "C2", "optionD": "D2",
            "correctOption": "B",
            "points": 10
        }
    ]

    res = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": batch_to_confirm, "publish": False})
    # Must fail (either 409 INTEGRITY_CONFLICT or 500)
    assert res.status_code in (409, 500)

    # CRITICAL EVIDENCE: Verify that row 1 was completely ROLLED BACK!
    conn = sqlite3.connect(db_file)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM questions WHERE code IN ('COMMITTED_OR_ROLLED_BACK_01', 'CRASH_ON_SECOND_ROW')")
    count_saved = cur.fetchone()[0]
    conn.close()
    assert count_saved == 0, f"Expected 0 questions saved due to transaction rollback, but found {count_saved}!"

def test_row_limit_2000_questions_and_2001_rejected(tmp_path):
    client, admin_h, student_h, db_file = setup_client(tmp_path)

    # 1. Confirm endpoint boundary
    questions_2000 = [
        {
            "code": f"Q_MAX_{i}",
            "subjectId": "physics",
            "grade": 10,
            "prompt": f"Câu hỏi tối đa thứ {i}",
            "optionA": "A", "optionB": "B", "optionC": "C", "optionD": "D",
            "correctOption": "A",
            "points": 10
        }
        for i in range(1, 2001)
    ]
    # Check 2001 questions in confirm payload -> rejected with 400 or 422
    questions_2001 = questions_2000 + [{
        "code": "Q_MAX_2001",
        "subjectId": "physics",
        "grade": 10,
        "prompt": "Câu hỏi thứ 2001",
        "optionA": "A", "optionB": "B", "optionC": "C", "optionD": "D",
        "correctOption": "A",
        "points": 10
    }]
    res_overflow = client.post("/admin/import/confirm", headers=admin_h, json={"validRows": questions_2001, "publish": False})
    assert res_overflow.status_code in (400, 422)

    # 2. CSV file with 2001 questions (2002 total rows with header)
    buf = io.StringIO()
    buf.write("Mã,Môn,Lớp,Chương,Chủ đề,Bộ sách,Câu hỏi,A,B,C,D,Đáp án đúng\n")
    for i in range(1, 2002):
        buf.write(f"QC_{i},physics,10,Ch1,T1,KNTT,Câu hỏi {i},A,B,C,D,A\n")
    csv_bytes = buf.getvalue().encode("utf-8")
    res_csv_over = client.post("/admin/import/preview", headers=admin_h, content=csv_bytes)
    assert res_csv_over.status_code == 200
    assert any("vượt quá số lượng câu hỏi tối đa" in e["message"] for e in res_csv_over.json()["errors"])
